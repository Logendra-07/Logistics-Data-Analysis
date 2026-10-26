
import os
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt

from docx import Document
from docx.shared import Inches

from sklearn.model_selection import train_test_split, GridSearchCV, KFold, cross_val_score
from sklearn.compose import ColumnTransformer
from sklearn.preprocessing import OneHotEncoder
from sklearn.pipeline import Pipeline
from sklearn.impute import SimpleImputer
from sklearn.ensemble import RandomForestRegressor
from sklearn.linear_model import LinearRegression
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score

# 1. Create output folders
os.makedirs("week4_output", exist_ok=True)
os.makedirs("output", exist_ok=True)

# 2. Load the Week 3 logistics dataset
data = pd.read_csv("data/logistics_data.csv")

print("Dataset shape:", data.shape)
print(data.head())

# 3. Select prediction target and useful features
target = "delivery_time_hours"
features = [
    "distance_km",
    "shipment_volume_kg",
    "transport_mode",
    "warehouse"
]

required = features + [target]
data = data.dropna(subset=[target]).copy()

X = data[features]
y = data[target]

numeric_features = ["distance_km", "shipment_volume_kg"]
categorical_features = ["transport_mode", "warehouse"]

# 4. Prepare numeric and categorical data
numeric_pipeline = Pipeline([
    ("imputer", SimpleImputer(strategy="median"))
])

categorical_pipeline = Pipeline([
    ("imputer", SimpleImputer(strategy="most_frequent")),
    ("encoder", OneHotEncoder(handle_unknown="ignore"))
])

preprocessor = ColumnTransformer([
    ("numeric", numeric_pipeline, numeric_features),
    ("categorical", categorical_pipeline, categorical_features)
])

# 5. Split data into training and testing sets
X_train, X_test, y_train, y_test = train_test_split(
    X, y, test_size=0.20, random_state=42
)

# 6. Train a baseline Linear Regression model
linear_model = Pipeline([
    ("preprocessor", preprocessor),
    ("model", LinearRegression())
])

linear_model.fit(X_train, y_train)
linear_predictions = linear_model.predict(X_test)

# 7. Train and tune a Random Forest model
forest_pipeline = Pipeline([
    ("preprocessor", preprocessor),
    ("model", RandomForestRegressor(random_state=42))
])

parameters = {
    "model__n_estimators": [50, 100],
    "model__max_depth": [None, 5, 10],
    "model__min_samples_leaf": [1, 2]
}

cv = KFold(n_splits=5, shuffle=True, random_state=42)

search = GridSearchCV(
    forest_pipeline,
    parameters,
    cv=cv,
    scoring="neg_mean_absolute_error",
    n_jobs=-1
)

search.fit(X_train, y_train)
best_model = search.best_estimator_
predictions = best_model.predict(X_test)

# 8. Evaluate both models
def evaluate(actual, predicted):
    return {
        "MAE": mean_absolute_error(actual, predicted),
        "RMSE": np.sqrt(mean_squared_error(actual, predicted)),
        "R2": r2_score(actual, predicted)
    }

results = pd.DataFrame({
    "Linear Regression": evaluate(y_test, linear_predictions),
    "Random Forest": evaluate(y_test, predictions)
}).T

results.to_csv("week4_output/model_evaluation.csv")
print("\nModel Evaluation:")
print(results)
print("\nBest Random Forest parameters:", search.best_params_)

# 9. Cross-validation on the training data
cv_scores = cross_val_score(
    best_model, X_train, y_train,
    cv=cv, scoring="neg_mean_absolute_error"
)

cv_mae = -cv_scores.mean()
print("Mean cross-validation MAE:", round(cv_mae, 3), "hours")

# 10. Save actual versus predicted results
comparison = pd.DataFrame({
    "Actual_Delivery_Time_Hours": y_test.values,
    "Predicted_Delivery_Time_Hours": predictions
})
comparison.to_csv("week4_output/predictions.csv", index=False)

# 11. Create actual versus predicted chart
plt.figure(figsize=(7, 5))
plt.scatter(y_test, predictions)
lower = min(y_test.min(), predictions.min())
upper = max(y_test.max(), predictions.max())
plt.plot([lower, upper], [lower, upper], linestyle="--")
plt.xlabel("Actual Delivery Time (hours)")
plt.ylabel("Predicted Delivery Time (hours)")
plt.title("Actual vs Predicted Delivery Time")
plt.tight_layout()
plt.savefig("week4_output/actual_vs_predicted.png", dpi=150)
plt.close()

# 12. Extract feature importance from the fitted model
encoder = (
    best_model.named_steps["preprocessor"]
    .named_transformers_["categorical"]
    .named_steps["encoder"]
)

feature_names = numeric_features + list(
    encoder.get_feature_names_out(categorical_features)
)

importances = best_model.named_steps["model"].feature_importances_

importance_table = pd.DataFrame({
    "Feature": feature_names,
    "Importance": importances
}).sort_values("Importance", ascending=False)

importance_table.to_csv(
    "week4_output/feature_importance.csv", index=False
)

plt.figure(figsize=(9, 5))
top_features = importance_table.head(10).sort_values("Importance")
plt.barh(top_features["Feature"], top_features["Importance"])
plt.xlabel("Feature Importance")
plt.title("Factors Influencing Predicted Delivery Time")
plt.tight_layout()
plt.savefig("week4_output/feature_importance.png", dpi=150)
plt.close()

# 13. Compare predicted delivery time by transport mode
scenario = X_test.copy()
scenario["transport_mode"] = scenario["transport_mode"].astype(str)

mode_results = []
for mode in sorted(data["transport_mode"].dropna().astype(str).unique()):
    test_scenario = scenario.copy()
    test_scenario["transport_mode"] = mode
    mode_results.append({
        "Transport_Mode": mode,
        "Average_Predicted_Hours": best_model.predict(
            test_scenario
        ).mean()
    })

mode_table = pd.DataFrame(mode_results).sort_values(
    "Average_Predicted_Hours"
)
mode_table.to_csv("week4_output/mode_optimization.csv", index=False)

# 14. Generate the detailed Word report
doc = Document()
doc.add_heading(
    "Week 4: Predictive Modeling and Optimization in Logistics Systems",
    0
)

doc.add_paragraph(
    "Project: Logistics Data Analysis\n"
    "Prediction target: Delivery time in hours\n"
    "Dataset: Week 3 simulated logistics dataset"
)

doc.add_heading("1. Objective", level=1)
doc.add_paragraph(
    "The objective is to predict shipment delivery time using shipment "
    "distance, shipment volume, transport mode and warehouse. The model "
    "results are used to propose data-driven logistics improvements."
)

doc.add_heading("2. Dataset and Problem Definition", level=1)
doc.add_paragraph(
    f"The dataset contains {len(data)} usable records and "
    f"{len(data.columns)} columns. The target variable is "
    "delivery_time_hours. The input variables are distance_km, "
    "shipment_volume_kg, transport_mode and warehouse. Delivery time "
    "is a regression target because it is a continuous numerical value."
)

doc.add_heading("3. Data Preparation", level=1)
doc.add_paragraph(
    "Rows without a target value were removed. Missing numerical input "
    "values are imputed using the median, while missing categorical "
    "values are imputed using the most frequent category. One-hot "
    "encoding converts categorical values into numerical model inputs. "
    "Preprocessing is fitted within the training pipeline to avoid "
    "using test data to learn preprocessing parameters."
)

doc.add_heading("4. Model Selection", level=1)
doc.add_paragraph(
    "Linear Regression is used as a baseline because it is simple and "
    "interpretable. Random Forest Regression is used to capture possible "
    "nonlinear relationships and interactions between logistics features. "
    "GridSearchCV tests multiple tree counts, maximum depths and minimum "
    "leaf sizes using five-fold cross-validation on the training data."
)

doc.add_heading("5. Training and Validation", level=1)
doc.add_paragraph(
    "The data is divided into 80% training and 20% testing sets using "
    "a fixed random seed. The test set is held out for final evaluation. "
    "The selected Random Forest model is also evaluated using five-fold "
    "cross-validation on the training set."
)

doc.add_heading("6. Performance Evaluation", level=1)
doc.add_paragraph(
    "Mean Absolute Error (MAE) measures average absolute prediction "
    "error in hours. Root Mean Squared Error (RMSE) gives greater "
    "weight to large errors. R-squared measures the proportion of "
    "target variation explained by the model; it can be negative when "
    "test predictions are worse than a constant-mean baseline."
)

table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
for cell, heading in zip(table.rows[0].cells,
                         ["Model", "MAE (hours)", "RMSE (hours)", "R-squared"]):
    cell.text = heading

for model_name, row in results.iterrows():
    cells = table.add_row().cells
    cells[0].text = model_name
    cells[1].text = f"{row['MAE']:.3f}"
    cells[2].text = f"{row['RMSE']:.3f}"
    cells[3].text = f"{row['R2']:.3f}"

doc.add_paragraph(
    f"Mean five-fold training cross-validation MAE for the tuned "
    f"Random Forest: {cv_mae:.3f} hours."
)
doc.add_paragraph(
    "The model with lower test MAE has the smaller average absolute "
    "error. Consider RMSE and R-squared alongside MAE before selecting "
    "a model. Results apply to this simulated dataset and are not a "
    "guarantee of real-world delivery performance."
)

doc.add_heading("7. Prediction and Feature Analysis", level=1)
doc.add_paragraph(
    "The actual-versus-predicted chart helps identify systematic "
    "prediction errors. Feature importance indicates which inputs "
    "contribute most to the fitted Random Forest predictions. Importance "
    "is not proof of causation and should be interpreted with care."
)
doc.add_picture("week4_output/actual_vs_predicted.png", width=Inches(6))
doc.add_picture("week4_output/feature_importance.png", width=Inches(6))

doc.add_heading("8. Optimization Strategies", level=1)
doc.add_paragraph(
    "1. Use predicted delivery times to flag shipments likely to miss "
    "their promised delivery window.\n"
    "2. Review the mode comparison table when assigning shipments. "
    "Treat it as a model-based scenario, not proof that one mode is "
    "always faster or cheaper.\n"
    "3. Prioritize dispatch planning and capacity allocation for "
    "high-risk shipments.\n"
    "4. Group compatible shipments and review route plans to reduce "
    "unnecessary distance and delays.\n"
    "5. Monitor actual delivery times, costs and prediction errors "
    "regularly, then retrain the model when new data becomes available."
)

doc.add_heading("9. Transport Mode Scenario Analysis", level=1)
doc.add_paragraph(
    "The following table predicts delivery time for the same test "
    "shipments while changing the transport mode. This isolates the "
    "model's scenario estimates while holding other listed features "
    "constant. It does not account for all real-world operational "
    "constraints or transport-mode costs."
)

mode_doc_table = doc.add_table(rows=1, cols=2)
mode_doc_table.style = "Table Grid"
mode_doc_table.rows[0].cells[0].text = "Transport Mode"
mode_doc_table.rows[0].cells[1].text = "Average Predicted Hours"

for _, row in mode_table.iterrows():
    cells = mode_doc_table.add_row().cells
    cells[0].text = str(row["Transport_Mode"])
    cells[1].text = f"{row['Average_Predicted_Hours']:.2f}"

doc.add_heading("10. Python Implementation", level=1)
doc.add_paragraph(
    "The complete implementation is available in "
    "scripts/week4_predictive_modeling.py. It loads data, prepares "
    "features, trains and tunes models, calculates evaluation metrics, "
    "creates charts, compares transport-mode scenarios and writes this "
    "report. The main libraries used are pandas, NumPy, scikit-learn, "
    "Matplotlib and python-docx."
)

doc.add_heading("11. Conclusion", level=1)
doc.add_paragraph(
    "This project demonstrates a complete predictive modeling workflow "
    "for logistics delivery-time forecasting. Model evaluation provides "
    "a quantitative basis for comparing approaches, while scenario "
    "analysis and feature importance help identify opportunities for "
    "better dispatch planning, capacity allocation and delivery "
    "monitoring. Future work should validate the model on larger "
    "real-world data and incorporate traffic, weather, service-level "
    "commitments and route constraints."
)

report_path = "week4_output/Week4_Predictive_Modeling_Report.docx"
doc.save(report_path)

print("\nWeek 4 completed successfully!")
print("Report:", report_path)
print("Evaluation: week4_output/model_evaluation.csv")
print("Predictions: week4_output/predictions.csv")
print("Feature importance: week4_output/feature_importance.csv")
print("Mode scenarios: week4_output/mode_optimization.csv")
print("Charts saved in week4_output/")
