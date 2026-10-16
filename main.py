
import os
import random
from datetime import datetime, timedelta

import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches

# Create folders
os.makedirs("data", exist_ok=True)
os.makedirs("output", exist_ok=True)

random.seed(42)

# ----------------------------------
# 1. Generate sample logistics data
# ----------------------------------

csv_file = "data/logistics_data.csv"

if not os.path.exists(csv_file):
    warehouses = ["Chennai", "Bengaluru", "Coimbatore"]
    destinations = ["Madurai", "Salem", "Trichy", "Vellore"]
    records = []

    for i in range(1, 101):
        promised = random.randint(2, 5)
        actual = max(
            1,
            promised + random.choice([-1, 0, 0, 0, 1, 2])
        )
        distance = random.randint(50, 500)
        cost = round(
            500 + distance * 4 + random.randint(0, 500), 2
        )

        records.append({
            "Order_ID": i,
            "Order_Date": (
                datetime(2026, 1, 1)
                + timedelta(days=random.randint(0, 89))
            ).strftime("%Y-%m-%d"),
            "Warehouse": random.choice(warehouses),
            "Destination": random.choice(destinations),
            "Distance_KM": distance,
            "Promised_Days": promised,
            "Delivery_Days": actual,
            "Shipping_Cost": cost,
            "Quantity": random.randint(1, 50)
        })

    pd.DataFrame(records).to_csv(csv_file, index=False)

# ----------------------------------
# 2. Load and clean data
# ----------------------------------

df = pd.read_csv(csv_file)

df.drop_duplicates(inplace=True)

numeric_columns = [
    "Distance_KM",
    "Promised_Days",
    "Delivery_Days",
    "Shipping_Cost",
    "Quantity"
]

for column in numeric_columns:
    df[column] = pd.to_numeric(df[column], errors="coerce")

df.dropna(subset=numeric_columns, inplace=True)

df = df[
    (df["Distance_KM"] > 0)
    & (df["Promised_Days"] > 0)
    & (df["Delivery_Days"] > 0)
    & (df["Shipping_Cost"] >= 0)
    & (df["Quantity"] > 0)
].copy()

if df.empty:
    raise ValueError("No valid logistics records found.")

# ----------------------------------
# 3. Calculate KPIs
# ----------------------------------

df["On_Time"] = (
    df["Delivery_Days"] <= df["Promised_Days"]
)

total_orders = len(df)
on_time_rate = df["On_Time"].mean() * 100
average_delivery = df["Delivery_Days"].mean()
average_cost = df["Shipping_Cost"].mean()
average_distance = df["Distance_KM"].mean()

warehouse_summary = df.groupby("Warehouse").agg(
    Orders=("Order_ID", "count"),
    Average_Cost=("Shipping_Cost", "mean"),
    On_Time_Rate=("On_Time", "mean")
).reset_index()

warehouse_summary["Average_Cost"] = (
    warehouse_summary["Average_Cost"].round(2)
)
warehouse_summary["On_Time_Rate"] = (
    warehouse_summary["On_Time_Rate"] * 100
).round(2)

# Save KPI results
warehouse_summary.to_csv(
    "output/warehouse_summary.csv", index=False
)

# ----------------------------------
# 4. Create chart
# ----------------------------------

chart_file = "output/warehouse_performance.png"

plt.figure(figsize=(8, 5))
plt.bar(
    warehouse_summary["Warehouse"],
    warehouse_summary["On_Time_Rate"]
)
plt.title("On-Time Delivery Rate by Warehouse")
plt.xlabel("Warehouse")
plt.ylabel("On-Time Delivery (%)")
plt.ylim(0, 100)
plt.tight_layout()
plt.savefig(chart_file, dpi=150)
plt.close()

# ----------------------------------
# 5. Create Word report
# ----------------------------------

doc = Document()

doc.add_heading(
    "Strategic Planning and Data Exploration in Logistics",
    0
)
doc.add_paragraph("Week 1 Task - Logistics Data Analysis")
doc.add_paragraph("Prepared by: Logendra Prabhu V")
doc.add_paragraph("Project: Logistics Performance Analysis")

doc.add_heading("1. Executive Summary", level=1)
doc.add_paragraph(
    "This project demonstrates a data-driven approach to "
    "improving logistics performance. Python is used to analyze "
    "sample delivery records, calculate key performance indicators, "
    "compare warehouses and visualize on-time delivery rates. "
    "The purpose is to demonstrate how data science can support "
    "delivery planning, cost monitoring and resource allocation. "
    "The dataset is synthetic and does not represent actual "
    "company operations."
)

doc.add_heading("2. Background Research", level=1)
doc.add_paragraph(
    "Logistics includes transportation, warehousing, inventory "
    "management and delivery scheduling. Delivery delays and "
    "high shipping costs can reduce customer satisfaction and "
    "operational efficiency. Data analysis can help organizations "
    "measure performance, detect patterns and support decisions. "
    "The World Bank Logistics Performance Index provides public "
    "context for studying logistics systems. Pandas supports "
    "data processing, Matplotlib supports visualization, and "
    "scikit-learn provides tools for future predictive modeling."
)
doc.add_paragraph(
    "Research resources:\n"
    "World Bank Logistics Performance Index: "
    "https://lpi.worldbank.org/\n"
    "Pandas documentation: https://pandas.pydata.org/docs/\n"
    "Scikit-learn documentation: https://scikit-learn.org/stable/\n"
    "Matplotlib documentation: https://matplotlib.org/stable/"
)

doc.add_heading("3. Project Definition", level=1)
doc.add_paragraph(
    "The scenario considers a logistics company operating "
    "warehouses in Chennai, Bengaluru and Coimbatore. It delivers "
    "orders to several destinations and wants to understand "
    "delivery reliability and shipping expenditure. The analysis "
    "compares warehouse performance and identifies opportunities "
    "for future delivery-time prediction and route optimization."
)

doc.add_heading("4. Project Objectives", level=1)
objectives = [
    "Measure the on-time delivery rate.",
    "Calculate average delivery time and shipping cost.",
    "Compare performance among warehouses.",
    "Explore patterns in delivery distance and delays.",
    "Propose predictive analytics and optimization methods."
]
for item in objectives:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("5. Key Performance Indicators", level=1)

kpis = [
    (
        "On-Time Delivery Rate",
        "Percentage of orders delivered on or before the promised time."
    ),
    (
        "Average Delivery Time",
        "Average number of days taken to deliver an order."
    ),
    (
        "Average Shipping Cost",
        "Average shipping expenditure per order."
    ),
    (
        "Average Delivery Distance",
        "Average distance travelled per order in kilometres."
    ),
    (
        "Total Orders",
        "Number of valid orders analyzed."
    )
]

table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
table.rows[0].cells[0].text = "KPI"
table.rows[0].cells[1].text = "Description"

for name, description in kpis:
    cells = table.add_row().cells
    cells[0].text = name
    cells[1].text = description

doc.add_heading("6. Literature and Data Research", level=1)
doc.add_paragraph(
    "Public logistics research and documentation can guide "
    "the design of this project. The World Bank Logistics "
    "Performance Index provides information about logistics "
    "performance across economies. For this demonstration, "
    "a synthetic dataset of 100 orders is generated with "
    "warehouse, destination, distance, promised delivery days, "
    "actual delivery days, shipping cost and quantity. This "
    "dataset allows the workflow to be tested without exposing "
    "private or confidential business information. Real-world "
    "implementation requires an appropriate public dataset or "
    "authorized company records."
)

doc.add_heading("7. Data Science Methodologies", level=1)

doc.add_heading("7.1 Exploratory Data Analysis", level=2)
doc.add_paragraph(
    "Descriptive statistics and charts can reveal patterns in "
    "delivery times, shipping costs and warehouse performance."
)

doc.add_heading("7.2 Regression", level=2)
doc.add_paragraph(
    "Regression can be used to estimate delivery time or shipping "
    "cost from variables such as distance, quantity, warehouse "
    "and historical delivery information. A model should be "
    "trained and evaluated using suitable historical data."
)

doc.add_heading("7.3 Clustering", level=2)
doc.add_paragraph(
    "Clustering can group similar orders or routes according "
    "to distance, cost and delivery performance. These groups "
    "may help identify routes that require additional attention."
)

doc.add_heading("7.4 Optimization", level=2)
doc.add_paragraph(
    "Optimization can compare feasible route assignments and "
    "delivery schedules while considering vehicle capacity, "
    "delivery windows, distance and cost constraints."
)

doc.add_heading("8. Strategic Roadmap", level=1)
roadmap = [
    "Define the logistics problem and project objectives.",
    "Collect order, delivery, warehouse and transportation data.",
    "Remove duplicates and handle missing or invalid values.",
    "Explore distributions, summary statistics and relationships.",
    "Calculate KPIs and compare warehouse performance.",
    "Develop and evaluate predictive models using historical data.",
    "Test route and scheduling optimization under real constraints.",
    "Present results and monitor performance after implementation."
]
for item in roadmap:
    doc.add_paragraph(item, style="List Number")

doc.add_heading("9. Python Code Illustration", level=1)
doc.add_paragraph(
    "The following snippet calculates the on-time delivery rate:"
)
doc.add_paragraph(
    'df = pd.read_csv("data/logistics_data.csv")\n'
    'df = df.drop_duplicates()\n'
    'df["On_Time"] = (\n'
    '    df["Delivery_Days"] <= df["Promised_Days"]\n'
    ')\n'
    'rate = df["On_Time"].mean() * 100\n'
    'print("On-time delivery rate:", round(rate, 2), "%")'
)

doc.add_heading("10. Analysis Results", level=1)
doc.add_paragraph(f"Valid orders analyzed: {total_orders}")
doc.add_paragraph(
    f"On-time delivery rate: {on_time_rate:.2f}%"
)
doc.add_paragraph(
    f"Average delivery time: {average_delivery:.2f} days"
)
doc.add_paragraph(
    f"Average shipping cost: {average_cost:.2f}"
)
doc.add_paragraph(
    f"Average delivery distance: {average_distance:.2f} km"
)
doc.add_paragraph(
    "These results are based on synthetic sample data and "
    "should not be interpreted as actual business performance."
)

doc.add_heading("11. Warehouse Performance Comparison", level=1)
doc.add_picture(chart_file, width=Inches(6))
doc.add_paragraph(
    warehouse_summary.to_string(index=False)
)

doc.add_heading("12. Expected Outcomes", level=1)
doc.add_paragraph(
    "The workflow is expected to improve visibility into "
    "delivery reliability, shipping costs and warehouse "
    "performance. Once real data is available, predictive "
    "models may help forecast delivery time and cost. Clustering "
    "may reveal similar operational patterns, while optimization "
    "may support better route assignments and scheduling. "
    "Any proposed improvements should be tested and evaluated "
    "before being applied in real operations."
)

doc.add_heading("13. Conclusion", level=1)
doc.add_paragraph(
    "This project demonstrates a practical Python-based "
    "workflow for strategic logistics analysis. It combines "
    "data cleaning, KPI calculation, visualization and a "
    "structured reporting process. The proposed roadmap provides "
    "a foundation for future regression, clustering and route "
    "optimization work. Further analysis with reliable real-world "
    "data is needed to validate findings and measure the effect "
    "of operational changes."
)

doc.add_heading("14. References", level=1)
references = [
    "World Bank, Logistics Performance Index: https://lpi.worldbank.org/",
    "Pandas documentation: https://pandas.pydata.org/docs/",
    "Scikit-learn documentation: https://scikit-learn.org/stable/",
    "Matplotlib documentation: https://matplotlib.org/stable/"
]
for reference in references:
    doc.add_paragraph(reference, style="List Bullet")

# Save Word report
report_file = "output/Logistics_Strategic_Report.docx"
doc.save(report_file)

print("\nAnalysis completed successfully!")
print("Total orders:", total_orders)
print("On-time delivery rate:", round(on_time_rate, 2), "%")
print("Average delivery time:", round(average_delivery, 2), "days")
print("Average shipping cost:", round(average_cost, 2))
print("Average delivery distance:", round(average_distance, 2), "km")
print("Word report:", report_file)
print("Performance chart:", chart_file)
print("Warehouse summary: output/warehouse_summary.csv")
print("Dataset:", csv_file)
