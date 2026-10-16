
import os
import pandas as pd
import numpy as np
from docx import Document
from docx.shared import Inches, Pt

# Create output folder
os.makedirs("week2_output", exist_ok=True)

# Load the Week 1 logistics dataset
df = pd.read_csv("data/logistics_data.csv")

# Add sample data quality issues for demonstration
df.loc[0, "Shipping_Cost"] = np.nan
df.loc[1, "Delivery_Days"] = np.nan
df = pd.concat([df, df.iloc[[2]]], ignore_index=True)

# Save the simulated raw dataset
df.to_csv("week2_output/raw_logistics_data.csv", index=False)

# Record initial dataset details
initial_rows = len(df)
missing_before = int(df.isnull().sum().sum())
duplicates_before = int(df.duplicated().sum())

# 1. Remove duplicate records
df = df.drop_duplicates().copy()

# 2. Handle missing numerical values using the median
numeric_columns = df.select_dtypes(include="number").columns
for column in numeric_columns:
    df[column] = df[column].fillna(df[column].median())

# 3. Detect outliers using the IQR method
outlier_counts = {}
for column in numeric_columns:
    q1 = df[column].quantile(0.25)
    q3 = df[column].quantile(0.75)
    iqr = q3 - q1
    lower = q1 - 1.5 * iqr
    upper = q3 + 1.5 * iqr
    outlier_counts[column] = int(
        ((df[column] < lower) | (df[column] > upper)).sum()
    )

# 4. Apply Min-Max normalization to selected numeric columns
normalization_columns = [
    column for column in
    ["Distance_KM", "Delivery_Days", "Shipping_Cost", "Quantity"]
    if column in df.columns
]

for column in normalization_columns:
    minimum = df[column].min()
    maximum = df[column].max()
    if maximum != minimum:
        df[column + "_Normalized"] = (
            (df[column] - minimum) / (maximum - minimum)
        )
    else:
        df[column + "_Normalized"] = 0.0

# Save cleaned data
df.to_csv("week2_output/cleaned_logistics_data.csv", index=False)

# Create a professional Word report
doc = Document()
doc.add_heading(
    "Week 2: Data Collection, Cleaning, and Preprocessing for Logistics Analysis",
    0
)

doc.add_paragraph("Project: Logistics Data Analysis")
doc.add_paragraph("Tools: Python, Pandas, NumPy, and python-docx")

doc.add_heading("1. Objective", level=1)
doc.add_paragraph(
    "The objective is to demonstrate a complete logistics data "
    "preprocessing pipeline. The workflow identifies missing values, "
    "removes duplicate records, detects potential numerical outliers, "
    "and normalizes selected features for future analysis."
)

doc.add_heading("2. Data Collection Simulation", level=1)
doc.add_paragraph(
    "The project uses the logistics CSV dataset created in Week 1. "
    "It contains sample order dates, warehouses, destinations, distances, "
    "promised and actual delivery times, shipping costs, and quantities. "
    "The dataset is synthetic and is used for demonstration, not as "
    "verified real-world logistics data. Missing values and a duplicate "
    "record are introduced to simulate common data collection problems."
)

doc.add_heading("3. Data Quality Assessment", level=1)
doc.add_paragraph(f"Rows before cleaning: {initial_rows}")
doc.add_paragraph(f"Missing cells before cleaning: {missing_before}")
doc.add_paragraph(f"Duplicate rows before cleaning: {duplicates_before}")

doc.add_heading("4. Cleaning Methodology", level=1)

doc.add_heading("4.1 Duplicate Records", level=2)
doc.add_paragraph(
    "Duplicate records may cause orders to be counted more than once. "
    "The pandas drop_duplicates() function removes identical rows. "
    "In a real system, order identifiers should also be checked to detect "
    "conflicting records with the same identifier."
)
doc.add_paragraph("Code: df = df.drop_duplicates().copy()")

doc.add_heading("4.2 Missing Values", level=2)
doc.add_paragraph(
    "Missing numerical values are filled using each column's median. "
    "The median is less sensitive to extreme values than the mean. "
    "In production, missing values should be investigated first, and "
    "categorical values may require a separate strategy."
)
doc.add_paragraph(
    "Code: df[column] = df[column].fillna(df[column].median())"
)

doc.add_heading("4.3 Outlier Detection", level=2)
doc.add_paragraph(
    "The Interquartile Range (IQR) method identifies observations below "
    "Q1 - 1.5 × IQR or above Q3 + 1.5 × IQR. Outliers are flagged for "
    "investigation rather than automatically deleted, because long routes "
    "and high shipping costs may be legitimate logistics events."
)
for column, count in outlier_counts.items():
    doc.add_paragraph(f"{column}: {count} potential outlier(s)")

doc.add_heading("4.4 Normalization", level=2)
doc.add_paragraph(
    "Min-Max normalization transforms selected numerical features to a "
    "range from 0 to 1. This can help algorithms that are sensitive to "
    "feature scales. Original columns are retained so the source values "
    "remain available for interpretation."
)
doc.add_paragraph(
    "Code: (df[column] - df[column].min()) / "
    "(df[column].max() - df[column].min())"
)

doc.add_heading("5. Results", level=1)
doc.add_paragraph(f"Rows after duplicate removal: {len(df)}")
doc.add_paragraph(
    f"Missing cells after numerical imputation: {int(df.isnull().sum().sum())}"
)
doc.add_paragraph(
    "The cleaned dataset and normalized features were saved to "
    "week2_output/cleaned_logistics_data.csv."
)

doc.add_heading("6. Impact on Logistics Analysis", level=1)
doc.add_paragraph(
    "High-quality data supports reliable delivery-time analysis, cost "
    "comparison, warehouse performance monitoring, and future predictive "
    "models. Duplicate orders can inflate totals, missing values can "
    "distort statistics, extreme observations can influence models, and "
    "inconsistent feature scales can affect some machine-learning methods. "
    "Preprocessing improves consistency, but each decision must be "
    "documented and validated against business knowledge."
)

doc.add_heading("7. Reflection and Limitations", level=1)
doc.add_paragraph(
    "This exercise demonstrates a reproducible preprocessing workflow. "
    "The data is synthetic, so the results do not establish actual "
    "logistics performance. The IQR method flags possible outliers but "
    "does not prove that they are errors. Before using real operational "
    "data, analysts should verify units, dates, identifiers, and business "
    "rules, and should avoid data leakage when preparing machine-learning "
    "datasets."
)

doc.add_heading("8. Conclusion", level=1)
doc.add_paragraph(
    "Data collection, cleaning, missing-value treatment, outlier "
    "assessment, and normalization are important steps before logistics "
    "data analysis. This project produces a cleaned CSV file and a report "
    "documenting the methods and their limitations."
)

doc.add_heading("9. References", level=1)
doc.add_paragraph("Pandas documentation: https://pandas.pydata.org/docs/")
doc.add_paragraph("NumPy documentation: https://numpy.org/doc/")
doc.add_paragraph(
    "Scikit-learn preprocessing guide: "
    "https://scikit-learn.org/stable/modules/preprocessing.html"
)

report_path = "week2_output/Week2_Logistics_Preprocessing_Report.docx"
doc.save(report_path)

print("\nWeek 2 preprocessing completed successfully!")
print("Initial rows:", initial_rows)
print("Rows after cleaning:", len(df))
print("Duplicate rows detected:", duplicates_before)
print("Missing cells before cleaning:", missing_before)
print("Missing cells after cleaning:", int(df.isnull().sum().sum()))
print("Report:", report_path)
print("Cleaned dataset: week2_output/cleaned_logistics_data.csv")
print("Raw simulated dataset: week2_output/raw_logistics_data.csv")
