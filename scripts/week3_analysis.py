
import os
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches

# Project folders
BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(BASE, "data")
OUTPUT = os.path.join(BASE, "output")

os.makedirs(DATA, exist_ok=True)
os.makedirs(OUTPUT, exist_ok=True)

np.random.seed(42)

# 1. Simulate a hypothetical logistics dataset
n = 200

df = pd.DataFrame({
    "shipment_id": range(1001, 1001 + n),
    "distance_km": np.random.randint(20, 1001, n),
    "shipment_volume_kg": np.random.randint(50, 2001, n),
    "transport_mode": np.random.choice(
        ["Truck", "Van", "Rail"], n, p=[0.55, 0.30, 0.15]
    ),
    "warehouse": np.random.choice(
        ["Chennai", "Bengaluru", "Hyderabad"], n
    )
})

# Simulate time and cost using distance, mode and random variation
mode_time = df["transport_mode"].map({
    "Truck": 0.045, "Van": 0.055, "Rail": 0.035
})
mode_cost = df["transport_mode"].map({
    "Truck": 12, "Van": 16, "Rail": 9
})

df["delivery_time_hours"] = (
    df["distance_km"] * mode_time
    + np.random.normal(5, 3, n)
).clip(lower=2).round(2)

df["transportation_cost"] = (
    df["distance_km"] * mode_cost
    + df["shipment_volume_kg"] * 1.5
    + np.random.normal(0, 500, n)
).clip(lower=100).round(2)

df["delay_hours"] = (
    np.random.exponential(3, n)
).round(2)

df["delivery_status"] = np.where(
    df["delay_hours"] > 5, "Delayed", "On Time"
)

# Save dataset
csv_path = os.path.join(DATA, "logistics_data.csv")
df.to_csv(csv_path, index=False)

# 2. Exploratory Data Analysis
numeric = [
    "distance_km",
    "shipment_volume_kg",
    "delivery_time_hours",
    "transportation_cost",
    "delay_hours"
]

summary = df[numeric].describe().round(2)
correlations = df[numeric].corr().round(2)
status_counts = df["delivery_status"].value_counts()
mode_summary = df.groupby("transport_mode")[
    ["delivery_time_hours", "transportation_cost"]
].mean().round(2)

print("\nDATASET PREVIEW")
print(df.head())

print("\nDATASET SHAPE:", df.shape)
print("\nMISSING VALUES")
print(df.isnull().sum())

print("\nSUMMARY STATISTICS")
print(summary)

print("\nCORRELATION MATRIX")
print(correlations)

print("\nDELIVERY STATUS")
print(status_counts)

print("\nAVERAGE TIME AND COST BY TRANSPORT MODE")
print(mode_summary)

# Save analysis tables
summary.to_csv(os.path.join(OUTPUT, "summary_statistics.csv"))
correlations.to_csv(os.path.join(OUTPUT, "correlation_matrix.csv"))
mode_summary.to_csv(os.path.join(OUTPUT, "transport_mode_summary.csv"))

# 3. Visualizations

# Chart 1: Distribution of delivery times
plt.figure(figsize=(8, 5))
plt.hist(df["delivery_time_hours"], bins=15, edgecolor="black")
plt.title("Distribution of Delivery Times")
plt.xlabel("Delivery Time (hours)")
plt.ylabel("Number of Shipments")
plt.tight_layout()
plt.savefig(os.path.join(OUTPUT, "delivery_time_distribution.png"))
plt.close()

# Chart 2: Average transportation cost by mode
plt.figure(figsize=(8, 5))
df.groupby("transport_mode")["transportation_cost"].mean().plot(
    kind="bar", edgecolor="black"
)
plt.title("Average Transportation Cost by Mode")
plt.xlabel("Transport Mode")
plt.ylabel("Average Cost (cost units)")
plt.xticks(rotation=0)
plt.tight_layout()
plt.savefig(os.path.join(OUTPUT, "shipping_cost_analysis.png"))
plt.close()

# Chart 3: Delivery status
plt.figure(figsize=(7, 5))
status_counts.plot(kind="bar", edgecolor="black")
plt.title("On-Time vs Delayed Shipments")
plt.xlabel("Delivery Status")
plt.ylabel("Number of Shipments")
plt.xticks(rotation=0)
plt.tight_layout()
plt.savefig(os.path.join(OUTPUT, "delivery_status.png"))
plt.close()

# Chart 4: Shipment volume versus transportation cost
plt.figure(figsize=(8, 5))
for mode, group in df.groupby("transport_mode"):
    plt.scatter(
        group["shipment_volume_kg"],
        group["transportation_cost"],
        label=mode,
        alpha=0.7
    )
plt.title("Shipment Volume vs Transportation Cost")
plt.xlabel("Shipment Volume (kg)")
plt.ylabel("Transportation Cost (cost units)")
plt.legend()
plt.tight_layout()
plt.savefig(os.path.join(OUTPUT, "volume_vs_cost.png"))
plt.close()

# Chart 5: Correlation heatmap using matplotlib
plt.figure(figsize=(8, 6))
plt.imshow(correlations, cmap="coolwarm", vmin=-1, vmax=1)
plt.colorbar(label="Correlation")
plt.xticks(range(len(numeric)), numeric, rotation=45, ha="right")
plt.yticks(range(len(numeric)), numeric)

for i in range(len(numeric)):
    for j in range(len(numeric)):
        plt.text(
            j, i, f"{correlations.iloc[i, j]:.2f}",
            ha="center", va="center", fontsize=8
        )

plt.title("Correlation Heatmap of Logistics Metrics")
plt.tight_layout()
plt.savefig(os.path.join(OUTPUT, "correlation_heatmap.png"))
plt.close()

# 4. Calculate key performance indicators
total = len(df)
on_time_pct = (df["delivery_status"] == "On Time").mean() * 100
delayed_pct = (df["delivery_status"] == "Delayed").mean() * 100
avg_time = df["delivery_time_hours"].mean()
avg_cost = df["transportation_cost"].mean()
avg_distance = df["distance_km"].mean()

# 5. Generate Word report
doc = Document()
doc.add_heading(
    "Week 3: Advanced Data Analysis and Visualization in Logistics",
    0
)

doc.add_paragraph(
    "Project: Logistics Data Analysis\n"
    "Method: Simulated data, exploratory analysis and visualization"
)

doc.add_heading("1. Introduction", level=1)
doc.add_paragraph(
    "This report examines a hypothetical logistics operation using Python. "
    "It analyzes shipment distance, shipment volume, delivery time, "
    "transportation cost, transport mode and delivery status. The purpose "
    "is to demonstrate how data analysis supports operational decisions."
)

doc.add_heading("2. Dataset Description", level=1)
doc.add_paragraph(
    f"The simulated dataset contains {total} shipments and "
    f"{len(df.columns)} columns. It includes shipment identifiers, distance "
    "in kilometres, shipment weight in kilograms, delivery time in hours, "
    "transportation cost in illustrative cost units, delay hours, transport "
    "mode and warehouse. The data is hypothetical and is not actual company data."
)

doc.add_heading("3. Methodology and Tools", level=1)
doc.add_paragraph(
    "NumPy generates reproducible sample data. Pandas performs data "
    "manipulation, descriptive statistics and correlation analysis. "
    "Matplotlib generates five visualizations. Python-docx creates this report."
)

doc.add_heading("4. Data Quality and Cleaning", level=1)
doc.add_paragraph(
    f"The dataset contains {int(df.isnull().sum().sum())} missing cells. "
    "The generated numeric values are constrained to plausible positive "
    "ranges. The dataset is checked for missing values and summary statistics. "
    "Since the dataset is generated in this script, no real-world records "
    "are claimed. In a real dataset, duplicate IDs, invalid values and "
    "outliers should be investigated before removal or correction."
)

doc.add_heading("5. Exploratory Data Analysis", level=1)
doc.add_paragraph("Summary statistics for the main numeric variables:")
table = doc.add_table(rows=1, cols=6)
table.style = "Light List Accent 1"
headers = ["Metric", "Mean", "Std. Dev.", "Minimum", "Median", "Maximum"]
for cell, value in zip(table.rows[0].cells, headers):
    cell.text = value

for col in numeric:
    row = table.add_row().cells
    values = [
        col,
        f"{df[col].mean():.2f}",
        f"{df[col].std():.2f}",
        f"{df[col].min():.2f}",
        f"{df[col].median():.2f}",
        f"{df[col].max():.2f}"
    ]
    for cell, value in zip(row, values):
        cell.text = value

doc.add_heading("6. Key Performance Indicators", level=1)
for item in [
    f"Total shipments analyzed: {total}",
    f"On-time delivery rate: {on_time_pct:.2f}%",
    f"Delayed shipment rate: {delayed_pct:.2f}%",
    f"Average delivery time: {avg_time:.2f} hours",
    f"Average transportation cost: {avg_cost:.2f} cost units",
    f"Average shipment distance: {avg_distance:.2f} km"
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("7. Visualizations and Interpretations", level=1)

charts = [
    (
        "delivery_time_distribution.png",
        "Figure 1: Delivery Time Distribution",
        "A histogram shows the spread and concentration of delivery times. "
        "Longer delivery times may indicate routes or shipments requiring "
        "further investigation."
    ),
    (
        "shipping_cost_analysis.png",
        "Figure 2: Average Transportation Cost by Mode",
        "A bar chart compares average cost across trucks, vans and rail. "
        "Differences may reflect distance, shipment mix and the simulated "
        "cost assumptions; they do not prove that one mode is always cheaper."
    ),
    (
        "delivery_status.png",
        "Figure 3: On-Time and Delayed Shipments",
        "A bar chart compares shipment counts by status. The on-time rate "
        "and delay rate help evaluate service reliability."
    ),
    (
        "volume_vs_cost.png",
        "Figure 4: Shipment Volume vs Transportation Cost",
        "A scatter plot explores the relationship between shipment weight "
        "and cost. An upward pattern would suggest that heavier shipments "
        "are associated with higher transportation costs."
    ),
    (
        "correlation_heatmap.png",
        "Figure 5: Correlation Between Logistics Metrics",
        "The heatmap summarizes linear relationships between numeric "
        "variables. Correlation does not establish causation."
    )
]

for filename, heading, explanation in charts:
    doc.add_heading(heading, level=2)
    doc.add_picture(os.path.join(OUTPUT, filename), width=Inches(6))
    doc.add_paragraph(explanation)

doc.add_heading("8. Findings and Recommendations", level=1)
doc.add_paragraph(
    f"The simulated dataset reports an on-time delivery rate of "
    f"{on_time_pct:.2f}% and an average delivery time of {avg_time:.2f} hours. "
    f"The average transportation cost is {avg_cost:.2f} cost units. "
    "These figures describe only the generated sample and should not be "
    "treated as industry benchmarks."
)

doc.add_paragraph(
    "Recommendations:\n"
    "1. Review delayed shipments by route, warehouse and transport mode.\n"
    "2. Compare transport modes using equivalent distances and shipment sizes.\n"
    "3. Monitor delivery time, cost per shipment and on-time delivery rate.\n"
    "4. Investigate extreme values before removing them from real data.\n"
    "5. Use actual historical records to validate findings before making "
    "operational changes.",
)

doc.add_heading("9. Limitations", level=1)
doc.add_paragraph(
    "The dataset is hypothetical, relatively small and generated using "
    "assumptions. Its patterns may not match real logistics operations. "
    "The results demonstrate analytical methods rather than predict actual "
    "business performance."
)

doc.add_heading("10. Conclusion", level=1)
doc.add_paragraph(
    "This project demonstrates a complete basic logistics analysis workflow: "
    "data simulation, quality checks, descriptive statistics, KPI calculation, "
    "correlation analysis, visual exploration and reporting. These techniques "
    "can help logistics teams identify questions for further investigation "
    "and support evidence-based planning when applied to reliable real data."
)

report_path = os.path.join(
    OUTPUT, "Week3_Logistics_Analysis_Report.docx"
)
doc.save(report_path)

print("\nSUCCESS!")
print("Dataset:", csv_path)
print("Report:", report_path)
print("Charts and analysis tables saved in:", OUTPUT)
